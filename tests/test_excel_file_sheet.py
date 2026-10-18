from datetime import date, datetime

import pytest
from openpyxl import load_workbook
from xlsx_helpers import DATA_DIR, mock_tabs, write_workbook

from shipment_tracker.display import StatusCell
from shipment_tracker.sheets.base import CellUpdate, TabNotFound
from shipment_tracker.sheets.excel_file_sheet import ExcelFileSheet, cell_text
from shipment_tracker.updater import update_statuses


@pytest.mark.parametrize(
    "value, text",
    [
        (None, ""),
        (123456789012, "123456789012"),
        (123456789012.0, "123456789012"),  # number-formatted tracking cell
        (2.5, "2.5"),
        (datetime(2026, 10, 6, 14, 30), "10/6/2026"),
        (date(2026, 10, 6), "10/6/2026"),
        ("1.00000E+11", "1.00000E+11"),  # already mangled text stays as-is
        (True, "TRUE"),
    ],
)
def test_cell_values_become_the_text_sheets_shows(value, text):
    assert cell_text(value) == text


def test_reads_tabs_from_a_sheets_style_download(tmp_path):
    path = write_workbook(mock_tabs(), tmp_path / "download.xlsx")
    rows = ExcelFileSheet(path).read_tab("Form Responses 1")
    assert rows[0][8] == "Tracking"
    assert rows[1][0] == "10/6/2026"  # stored as a date in the file
    assert rows[1][8] == "100000000001"  # stored as a number in the file


def test_unknown_tab_lists_the_tabs_that_exist(tmp_path):
    path = write_workbook(mock_tabs(), tmp_path / "download.xlsx")
    with pytest.raises(TabNotFound, match="Manual Entries"):
        ExcelFileSheet(path).read_tab("Typo")


def test_saves_text_and_fill_to_a_new_file_and_leaves_the_download_alone(tmp_path):
    path = write_workbook(mock_tabs(), tmp_path / "download.xlsx")
    original_bytes = path.read_bytes()

    sheet = ExcelFileSheet(path)
    update_statuses(sheet, "Form Responses 1", {"100000000001": StatusCell("Delivered 10/14", "#b7e1cd")})
    output = sheet.save(tmp_path / "out" / "with-status.xlsx")

    assert path.read_bytes() == original_bytes
    saved = load_workbook(output)["Form Responses 1"]
    assert saved["J1"].value == "Status"
    assert saved["J2"].value == "Delivered 10/14"
    assert saved["J2"].fill.start_color.rgb.endswith("B7E1CD")
    assert saved["B2"].value == "Avery Example"  # human columns untouched


def test_clearing_a_color_removes_the_fill(tmp_path):
    path = write_workbook({"T": [["Tracking"], ["111111111111"]]}, tmp_path / "d.xlsx")
    sheet = ExcelFileSheet(path)
    sheet.write_cells("T", [CellUpdate(2, 2, "x", "#f4c7c3")])
    sheet.write_cells("T", [CellUpdate(2, 2, "y", None)])
    saved = load_workbook(sheet.save(tmp_path / "o.xlsx"))["T"]
    assert saved["B2"].fill.fill_type is None


def test_refuses_to_overwrite_the_downloaded_file(tmp_path):
    path = write_workbook(mock_tabs(), tmp_path / "download.xlsx")
    with pytest.raises(ValueError, match="overwrite"):
        ExcelFileSheet(path).save(path)


def test_committed_sample_workbook_matches_the_mock_sheet():
    sheet = ExcelFileSheet(DATA_DIR / "sample_tracking.xlsx")
    for tab, rows in mock_tabs().items():
        assert sheet.read_tab(tab) == rows
