"""A spreadsheet backend for an Excel file downloaded from Google Sheets.

Until the tool can talk to Google Sheets directly, the workflow is:

1. In Google Sheets: File > Download > Microsoft Excel (.xlsx). One file
   holds every tab.
2. The tool reads that file and writes Status cells into memory.
3. `save()` writes a NEW file. The downloaded file is never changed.
4. You copy the Status column from the new file back into Google Sheets.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from shipment_tracker.sheets.base import CellUpdate, SheetClient, TabNotFound


class ExcelFileSheet(SheetClient):
    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        # data_only=True gives a formula's last calculated value, not the formula text.
        self._workbook = load_workbook(self.path, data_only=True)

    def read_tab(self, tab: str) -> list[list[str]]:
        sheet = self._sheet(tab)
        return [[cell_text(value) for value in row] for row in sheet.iter_rows(values_only=True)]

    def write_cells(self, tab: str, updates: list[CellUpdate]) -> None:
        sheet = self._sheet(tab)
        for update in updates:
            cell = sheet.cell(row=update.row, column=update.column)
            cell.value = update.text
            if update.color is None:
                cell.fill = PatternFill(fill_type=None)
            else:
                rgb = update.color.lstrip("#").upper()
                cell.fill = PatternFill(fill_type="solid", start_color=rgb, end_color=rgb)

    def save(self, output_path: str | Path) -> Path:
        """Write the updated workbook to a new file. Refuses to overwrite the input."""
        output_path = Path(output_path)
        if output_path.resolve() == self.path.resolve():
            raise ValueError("Refusing to overwrite the downloaded file. Choose a different output path.")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        self._workbook.save(output_path)
        return output_path

    def _sheet(self, tab: str):
        if tab not in self._workbook.sheetnames:
            raise TabNotFound(f"No tab named '{tab}'. Tabs in the file: {self._workbook.sheetnames}")
        return self._workbook[tab]


def cell_text(value) -> str:
    """Turn an Excel cell value into the text you'd see in Google Sheets.

    Matters most for tracking numbers: Sheets often stores them as numbers,
    and Excel files keep the full value (123456789012.0), so we print it
    back as plain digits instead of "1.23457E+11".
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return f"{value.month}/{value.day}/{value.year}"
    if isinstance(value, date):
        return f"{value.month}/{value.day}/{value.year}"
    return str(value)
